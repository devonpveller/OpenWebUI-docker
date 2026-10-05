"""Tests for frontend/owui/tools/pantry.py - no OWUI, no live service.

A stdlib http.server bound to 127.0.0.1 stands in for openbrain-pantry and records
method, path, headers and body. A socket guard fails the run if anything connects
anywhere but 127.0.0.1.

Run (from the repo root; xlsx tests need openpyxl, otherwise they are SKIPPED):

    python -m unittest discover -s frontend/owui/tools/tests -v

Full run with openpyxl, in a throwaway container (the production OWUI image already has
openpyxl + pydantic), no network, labelled:

    docker run --rm --network none --label ai-stack.harness.owner=<id> \
      -v "<repo>/frontend/owui:/o:ro" -e PYTHONDONTWRITEBYTECODE=1 -e PANTRY_REQUIRE_OWUI=1 \
      --entrypoint python openwebui:local -m unittest discover -s /o/tools/tests -v
"""

import asyncio
import csv
import http.server
import importlib.util
import io
import json
import os
import socket
import sys
import tempfile
import threading
import unittest

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

HERE = os.path.dirname(os.path.abspath(__file__))
_spec = importlib.util.spec_from_file_location("pantry_tool", os.path.join(HERE, "..", "pantry.py"))
pantry = importlib.util.module_from_spec(_spec)
sys.modules["pantry_tool"] = pantry
_spec.loader.exec_module(pantry)

_real_connect = socket.socket.connect
_real_connect_ex = socket.socket.connect_ex
BAD_CONNECTS = []


def _host_of(addr):
    return addr[0] if isinstance(addr, tuple) else str(addr)


def _guarded_connect(self, addr):
    if self.family in (socket.AF_INET, socket.AF_INET6) and _host_of(addr) != "127.0.0.1":
        BAD_CONNECTS.append(addr)
        raise OSError(f"test guard: connect to {addr} refused (only 127.0.0.1 allowed)")
    return _real_connect(self, addr)


def _guarded_connect_ex(self, addr):
    if self.family in (socket.AF_INET, socket.AF_INET6) and _host_of(addr) != "127.0.0.1":
        BAD_CONNECTS.append(addr)
        return 111
    return _real_connect_ex(self, addr)


class FakeService:
    """Records every request; replies from `routes[(method, path)]` = (status, body)."""

    def __init__(self):
        self.requests = []
        self.routes = {}
        outer = self

        class H(http.server.BaseHTTPRequestHandler):
            def log_message(self, *a):
                pass

            def _handle(self):
                n = int(self.headers.get("Content-Length") or 0)
                raw = self.rfile.read(n) if n else b""
                path, _, query = self.path.partition("?")
                rec = {
                    "method": self.command,
                    "path": path,
                    "query": query,
                    "headers": {k.lower(): v for k, v in self.headers.items()},
                    "body": json.loads(raw) if raw else None,
                }
                outer.requests.append(rec)
                spec = outer.routes.get((self.command, path), (200, {}))
                status, body = spec(rec) if callable(spec) else spec
                data = json.dumps(body).encode()
                self.send_response(status)
                self.send_header("Content-Type", "application/json")
                self.send_header("Content-Length", str(len(data)))
                self.end_headers()
                self.wfile.write(data)

            do_GET = do_POST = do_PUT = _handle

        self.httpd = http.server.ThreadingHTTPServer(("127.0.0.1", 0), H)
        self.port = self.httpd.server_address[1]
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()

    def stop(self):
        self.httpd.shutdown()
        self.httpd.server_close()


def run(coro):
    return asyncio.run(coro)


STOCK = [
    {"id": "i1", "name": "Chickpeas, canned", "aliases": [], "category": "Canned", "kind": "counted",
     "quantity": 4, "unit": "count", "level": None, "location": "pantry", "expires_on": "2027-01-31",
     "allergens": [], "may_contain": [], "reserved": 1, "available": 3, "use_soon": False},
    {"id": "i2", "name": "Olive oil", "aliases": [], "category": "Oils", "kind": "staple",
     "quantity": None, "unit": None, "level": "low", "location": "cupboard", "expires_on": None,
     "allergens": [], "may_contain": [], "reserved": 0, "available": None, "use_soon": False},
    {"id": "i3", "name": "Pesto", "aliases": [], "category": "Sauces", "kind": "counted",
     "quantity": 190.5, "unit": "g", "level": None, "location": "fridge", "expires_on": "2026-10-20",
     "allergens": ["tree nuts", "milk"], "may_contain": [], "reserved": 0, "available": 190.5, "use_soon": True},
]


class Base(unittest.TestCase):
    def setUp(self):
        BAD_CONNECTS.clear()
        socket.socket.connect = _guarded_connect
        socket.socket.connect_ex = _guarded_connect_ex
        self.svc = FakeService()
        self.addCleanup(self.svc.stop)
        self.tool = pantry.Tools()
        self.tool.valves.service_url = f"http://127.0.0.1:{self.svc.port}"
        self.tool.valves.api_key = "k-test"
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)

    def tearDown(self):
        socket.socket.connect = _real_connect
        socket.socket.connect_ex = _real_connect_ex
        self.assertEqual(BAD_CONNECTS, [], "a socket tried to leave 127.0.0.1")

    def last(self):
        return self.svc.requests[-1]

    def attach(self, name, data: bytes):
        p = os.path.join(self.tmp.name, name)
        with open(p, "wb") as f:
            f.write(data)
        return [{"type": "file", "id": "f1", "name": name, "file": {"id": "f1", "filename": name, "path": p}}]

    def preview_echo(self):
        def h(rec):
            rows = rec["body"]["rows"]
            return 200, {"preview_id": "pv1", "changed": [], "new": [], "missing": [], "unmatched": [],
                         "unconvertible": [], "echo_row_count": len(rows)}
        self.svc.routes[("POST", "/audit/preview")] = h


class Loading(unittest.TestCase):
    def test_instantiates_outside_owui(self):
        t = pantry.Tools()
        self.assertEqual(t.valves.service_url, "http://openbrain-pantry:8000")
        self.assertEqual(t.valves.api_key, "")
        self.assertGreater(t.valves.request_timeout_s, 0)

    def test_only_documented_functions_and_routes(self):
        expected = {
            "get_pantry", "update_pantry", "get_guidance", "save_recipe", "save_household_recipe",
            "cook", "log_cooked_meal", "correct_cook", "plan_meal", "get_plan", "set_plan_status", "shopping_list",
            "restock", "record_evaluation", "propose_preferences", "confirm_preferences",
            "manage_household", "export_pantry", "import_pantry", "accuracy_report",
            "confirm_import",
        }
        public = {n for n in dir(pantry.Tools) if not n.startswith("_") and callable(getattr(pantry.Tools, n)) and n != "Valves"}
        self.assertEqual(public, expected)
        for n in expected:
            self.assertTrue((getattr(pantry.Tools, n).__doc__ or "").strip(), n)
        # No stock arithmetic and no undocumented route: every route literal in the source is in API.md.
        with open(os.path.join(HERE, "..", "pantry.py"), encoding="utf-8") as fh:
            src = fh.read()
        import re
        routes = set(re.findall(r'"(/(?:pantry|settings|people|recipes|cook|plan|shopping-list|restock|audit|accuracy|evaluations|preferences|hypotheses|explored|guidance)[^"?]*)"', src))
        routes = {re.sub(r"\{[^}]*\}", "{id}", r) for r in routes}
        documented = {"/pantry", "/pantry/adjust", "/guidance", "/recipes", "/cook", "/cook/{id}/correct", "/plan", "/plan/{id}/status", "/shopping-list",
                      "/restock", "/evaluations", "/preferences", "/preferences/confirm", "/people", "/settings",
                      "/audit/preview", "/audit/commit", "/accuracy"}
        self.assertTrue(routes <= documented, routes - documented)


class Functions(Base):
    def test_get_pantry(self):
        self.svc.routes[("GET", "/pantry")] = (200, {"items": STOCK})
        out = json.loads(run(self.tool.get_pantry(q="oil", category="Oils", use_soon=True)))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("GET", "/pantry"))
        self.assertEqual(r["headers"]["x-pantry-key"], "k-test")
        for part in ("q=oil", "category=Oils", "use_soon=true"):
            self.assertIn(part, r["query"])
        self.assertEqual(out["total"], 3)

    def test_get_pantry_truncates(self):
        self.tool.valves.list_limit = 2
        self.svc.routes[("GET", "/pantry")] = (200, {"items": STOCK})
        out = json.loads(run(self.tool.get_pantry()))
        self.assertTrue(out["truncated"])
        self.assertEqual(len(out["items"]), 2)

    def test_update_pantry(self):
        self.svc.routes[("POST", "/pantry/adjust")] = (200, {"applied": [], "created": [], "unmatched": [{"name": "x", "candidates": []}], "unconvertible": []})
        out = json.loads(run(self.tool.update_pantry([{"name": "rice", "delta": 500, "unit": "g", "create": True}], reason="manual")))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/pantry/adjust"))
        self.assertEqual(r["body"], {"reason": "manual", "items": [{"name": "rice", "delta": 500, "unit": "g", "create": True}]})
        self.assertTrue(any("unmatched" in a for a in out["attention"]))

    def test_update_pantry_accepts_json_string(self):
        run(self.tool.update_pantry('[{"id":"i1","quantity":2}]', reason="correct"))
        self.assertEqual(self.last()["body"]["items"], [{"id": "i1", "quantity": 2}])

    def test_get_guidance(self):
        run(self.tool.get_guidance(theme="soup", guest_context={"adults": 2, "allergies": ["nuts"]}))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/guidance"))
        self.assertEqual(r["body"], {"theme": "soup", "guest_context": {"adults": 2, "allergies": ["nuts"]}})

    def test_save_recipe_and_household(self):
        ing = [{"pantry_item_id": "i1", "name": "chickpeas", "quantity": 2, "unit": "count"}]
        run(self.tool.save_recipe("Curry", 4, ing, ["cook"], theme="curry", cuisine="indian", draft=True))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/recipes"))
        self.assertEqual(r["body"]["source"], "generated")
        self.assertEqual(r["body"]["servings"], 4)
        self.assertEqual(r["body"]["draft"], True)
        self.assertNotIn("id", r["body"])
        run(self.tool.save_recipe("Curry", 4, ing, ["cook"], recipe_id="r9", reason="restock swap"))
        self.assertEqual(self.last()["body"]["id"], "r9")
        run(self.tool.save_household_recipe("Our chili", 6, ing, ["stew"], rotation=True))
        b = self.last()["body"]
        self.assertEqual((b["source"], b["rotation"]), ("household", True))

    def test_cook_and_log(self):
        self.svc.routes[("POST", "/cook")] = (201, {"cook_event_id": "c1", "deductions": [], "shortfalls": [{"name": "x"}], "unconvertible": []})
        out = json.loads(run(self.tool.cook("r1", servings=3, meal_plan_id="p1")))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/cook"))
        self.assertEqual(r["body"], {"recipe_id": "r1", "servings": 3, "meal_plan_id": "p1"})
        self.assertTrue(any("shortfalls" in a for a in out["attention"]))
        run(self.tool.log_cooked_meal("r2", cooked_at="2026-10-03"))
        self.assertEqual(self.last()["body"], {"recipe_id": "r2", "logged_after": True, "cooked_at": "2026-10-03"})

    def test_correct_cook(self):
        run(self.tool.correct_cook("c1", adjustments=[{"item_id": "i2", "actual_used": 30, "unit": "ml"}]))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/cook/c1/correct"))
        self.assertEqual(r["body"], {"adjustments": [{"item_id": "i2", "actual_used": 30, "unit": "ml"}]})
        run(self.tool.correct_cook("c1", undo=True))
        self.assertEqual(self.last()["body"], {"undo": True})

    def test_plan_and_get_plan(self):
        run(self.tool.plan_meal("2026-10-06", recipe_id="r1", guest_context={"adults": 2}))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/plan"))
        self.assertEqual(r["body"], {"date": "2026-10-06", "recipe_id": "r1", "guest_context": {"adults": 2}})
        run(self.tool.get_plan(week_start="2026-10-05"))
        r = self.last()
        self.assertEqual((r["method"], r["path"], r["query"]), ("GET", "/plan", "week_start=2026-10-05"))

    def test_set_plan_status(self):
        run(self.tool.set_plan_status("p7", "skipped"))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/plan/p7/status"))
        self.assertEqual(r["body"], {"status": "skipped"})
        self.assertEqual(r["headers"]["x-pantry-key"], "k-test")

    def test_shopping_and_restock(self):
        run(self.tool.shopping_list("2026-10-05"))
        r = self.last()
        self.assertEqual((r["method"], r["path"], r["body"]), ("POST", "/shopping-list", {"week_start": "2026-10-05"}))
        run(self.tool.restock("l1", bought="all", except_items=["leeks"], actual=[{"name": "flour", "quantity": 1, "unit": "kg"}]))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/restock"))
        self.assertEqual(r["body"], {"list_id": "l1", "bought": "all", "except": ["leeks"], "actual": [{"name": "flour", "quantity": 1, "unit": "kg"}]})

    def test_taste_functions(self):
        run(self.tool.record_evaluation("c1", who="child", liked=False, why="slimy", exposures=[{"subject": "mushroom", "reaction": "refused"}]))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/evaluations"))
        self.assertEqual(r["body"]["cook_event_id"], "c1")
        self.assertEqual(r["body"]["who"], "child")
        self.assertEqual(r["body"]["exposures"][0]["reaction"], "refused")
        st = [{"statement": "no raw mushroom", "strength": "contextual", "subject": "mushroom", "scope": "recipe", "who": "adult"}]
        run(self.tool.propose_preferences(st))
        self.assertEqual((self.last()["method"], self.last()["path"], self.last()["body"]), ("POST", "/preferences", {"statements": st}))
        run(self.tool.confirm_preferences(["p1"], edits={"p1": {"strength": "soft"}}, reject=["p2"]))
        r = self.last()
        self.assertEqual((r["path"], r["body"]), ("/preferences/confirm", {"ids": ["p1"], "edits": {"p1": {"strength": "soft"}}, "reject": ["p2"]}))

    def test_household(self):
        run(self.tool.manage_household("list"))
        self.assertEqual((self.last()["method"], self.last()["path"]), ("GET", "/people"))
        run(self.tool.manage_household("save", label="Kid", role="child", birth_month="2022-04", allergies=["egg"]))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/people"))
        self.assertEqual(r["body"], {"label": "Kid", "role": "child", "birth_month": "2022-04", "allergies": ["egg"]})
        run(self.tool.manage_household("get_settings"))
        self.assertEqual((self.last()["method"], self.last()["path"]), ("GET", "/settings"))
        run(self.tool.manage_household("set_settings", settings={"child_cooldown_days": 10}))
        r = self.last()
        self.assertEqual((r["method"], r["path"], r["body"]), ("PUT", "/settings", {"child_cooldown_days": 10}))
        bad = json.loads(run(self.tool.manage_household("nonsense")))
        self.assertEqual(bad["error"], "invalid")

    def test_accuracy(self):
        run(self.tool.accuracy_report(since="2026-09-01"))
        r = self.last()
        self.assertEqual((r["method"], r["path"], r["query"]), ("GET", "/accuracy", "since=2026-09-01"))

    def test_confirm_import(self):
        run(self.tool.confirm_import("pv1", missing={"i9": "zero"}, exclude=["i1"]))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/audit/commit"))
        self.assertEqual(r["body"], {"preview_id": "pv1", "missing": {"i9": "zero"}, "exclude": ["i1"]})


# Row-carrying parameters: what the model must send as real objects. OWUI 0.11 renders a
# bare `list` as {"type":"array","items":{"type":"string"}}, so the model sends each row
# as a STRING (2026-10-04 live failure: update_pantry -> items[0] must be an object).
OBJECT_ARRAYS = {
    ("update_pantry", "items"),
    ("save_recipe", "ingredients"),
    ("save_household_recipe", "ingredients"),
    ("correct_cook", "adjustments"),
    ("restock", "substitutions"),
    ("restock", "actual"),
    ("record_evaluation", "exposures"),
    ("propose_preferences", "statements"),
}
OBJECTS = {
    ("get_guidance", "guest_context"),
    ("save_recipe", "guest_context"),
    ("cook", "guest_context"),
    ("log_cooked_meal", "guest_context"),
    ("plan_meal", "guest_context"),
    ("confirm_preferences", "edits"),
    ("manage_household", "settings"),
    ("confirm_import", "missing"),
}
STRING_ARRAYS = {
    ("save_recipe", "instructions"),
    ("save_recipe", "tags"),
    ("save_household_recipe", "instructions"),
    ("save_household_recipe", "tags"),
    ("restock", "except_items"),
    ("restock", "bought"),
    ("confirm_preferences", "ids"),
    ("confirm_preferences", "reject"),
    ("manage_household", "allergies"),
    ("confirm_import", "exclude"),
}

_owui_get_tool_specs = None
if importlib.util.find_spec("open_webui") is not None:  # only inside the openwebui image
    # Importing OWUI initialises its env/db: give it a throwaway secret and data dir.
    if not os.environ.get("WEBUI_SECRET_KEY"):
        os.environ["WEBUI_SECRET_KEY"] = "test-only-not-a-secret"
    os.environ["DATA_DIR"] = tempfile.mkdtemp(prefix="owui-data-")
    try:
        from open_webui.utils.tools import get_tool_specs as _owui_get_tool_specs
    except (Exception, SystemExit) as _e:  # pragma: no cover
        print("open_webui present but not importable:", repr(_e))
        _owui_get_tool_specs = None


def _unwrap(prop):
    """Drop the null branch of an Optional[...] schema."""
    if isinstance(prop, dict) and "anyOf" in prop:
        rest = [b for b in prop["anyOf"] if b.get("type") != "null"]
        if len(rest) == 1:
            return rest[0]
    return prop


def _strip_optional(ann):
    import typing

    args = typing.get_args(ann)
    if type(None) in args:
        return [a for a in args if a is not type(None)][0]
    return ann


class ToolSchema(unittest.TestCase):
    """The parameter schema the model sees. Runs OWUI's real get_tool_specs when
    open_webui is importable (container run); otherwise checks the annotations."""

    def _check(self, kind_set, check):
        import typing

        if _owui_get_tool_specs is not None:
            specs = _owui_get_tool_specs(pantry.Tools())
            props = {s["name"]: s["parameters"].get("properties", {}) for s in specs}
            for fn, param in sorted(kind_set):
                with self.subTest(fn=fn, param=param, via="get_tool_specs"):
                    check(_unwrap(props[fn][param]), f"{fn}.{param}")
        else:
            for fn, param in sorted(kind_set):
                with self.subTest(fn=fn, param=param, via="annotations"):
                    ann = typing.get_type_hints(getattr(pantry.Tools, fn))[param]
                    check(_strip_optional(ann), f"{fn}.{param}")

    def test_restock_descriptions(self):
        """The model reads these; checked in the real get_tool_specs output when available."""
        if _owui_get_tool_specs is not None:
            specs = {s["name"]: s["parameters"]["properties"] for s in _owui_get_tool_specs(pantry.Tools())}
            ex = specs["restock"]["except_items"]["description"]
            bo = specs["restock"]["bought"]["description"]
        else:
            ex = bo = pantry.Tools.restock.__doc__
        self.assertNotIn('bought="all"', ex)
        self.assertIn("omit bought", ex.lower())
        self.assertIn("empty list", bo.lower())
        self.assertIn("nothing was bought", bo.lower())

    def test_which_path_ran(self):
        if os.environ.get("PANTRY_REQUIRE_OWUI"):  # container run: no silent fallback
            self.assertIsNotNone(_owui_get_tool_specs, "PANTRY_REQUIRE_OWUI set but get_tool_specs did not import")
        print("\n[ToolSchema] path:", "OWUI get_tool_specs" if _owui_get_tool_specs else "annotations only (no open_webui)")

    def test_object_arrays(self):
        import typing

        def chk(x, name):
            if isinstance(x, dict):
                self.assertEqual(x.get("type"), "array", name)
                self.assertEqual(x["items"].get("type"), "object", name)
            else:
                self.assertIs(typing.get_origin(x), list, name)
                self.assertEqual(typing.get_args(x), (dict,), name)

        self._check(OBJECT_ARRAYS, chk)

    def test_objects(self):
        def chk(x, name):
            if isinstance(x, dict):
                self.assertEqual(x.get("type"), "object", name)
            else:
                self.assertIs(x, dict, name)

        self._check(OBJECTS, chk)

    def test_string_arrays_stay_strings(self):
        import typing

        def chk(x, name):
            if isinstance(x, dict):
                self.assertEqual(x.get("type"), "array", name)
                self.assertEqual(x["items"].get("type"), "string", name)
            else:
                self.assertIs(typing.get_origin(x) or x, list, name)
                self.assertNotIn(dict, typing.get_args(x), name)

        self._check(STRING_ARRAYS, chk)

    def test_every_list_or_dict_param_is_classified(self):
        """A new list/dict parameter must be added to one of the sets above."""
        import inspect
        import typing

        known = OBJECT_ARRAYS | OBJECTS | STRING_ARRAYS
        # not model-visible (OWUI-injected) or documented as a JSON string / union
        skip = {("import_pantry", "mapping")}
        for fn in pantry.Tools.__dict__:
            if fn.startswith("_") or not inspect.iscoroutinefunction(getattr(pantry.Tools, fn)):
                continue
            for param, ann in typing.get_type_hints(getattr(pantry.Tools, fn)).items():
                if param.startswith("__") or param == "return" or (fn, param) in skip:
                    continue
                base = _strip_optional(ann)
                if (typing.get_origin(base) or base) in (list, dict):
                    self.assertIn((fn, param), known, f"{fn}.{param} is a list/dict param not classified")


class StringElementRows(Base):
    """2026-10-04 live failure: the model sent each row as a JSON-object STRING."""

    def test_update_pantry_json_object_strings_become_objects(self):
        rows = ['{"name":"Flour","kind":"staple","level":"plenty","create":true}',
                ' {"name":"Rice","delta":500,"unit":"g","create":true} ']
        run(self.tool.update_pantry(items=rows))
        body = self.last()["body"]
        self.assertEqual(body["items"][0], {"name": "Flour", "kind": "staple", "level": "plenty", "create": True})
        self.assertEqual(body["items"][1]["name"], "Rice")
        self.assertTrue(all(isinstance(i, dict) for i in body["items"]))

    def test_non_json_string_element_sends_nothing(self):
        out = json.loads(run(self.tool.update_pantry(items=['{"name":"Flour"}', "Flour, 2 bags"])))
        self.assertEqual(self.svc.requests, [])
        self.assertFalse(out["ok"])
        self.assertEqual(out["error"], "invalid_row")
        self.assertIn("items[1]", out["detail"])
        self.assertIn("Flour, 2 bags", out["detail"])

    def test_json_array_element_string_is_not_an_object(self):
        out = json.loads(run(self.tool.update_pantry(items=["[1,2]"])))
        self.assertEqual(self.svc.requests, [])
        self.assertEqual(out["error"], "invalid_row")

    def test_every_row_param_parses_string_elements(self):
        row = '{"name":"x","quantity":1,"unit":"g"}'
        run(self.tool.save_recipe(name="r", servings=2, ingredients=[row], instructions=["a"]))
        self.assertEqual(self.last()["body"]["ingredients"], [{"name": "x", "quantity": 1, "unit": "g"}])
        run(self.tool.save_household_recipe(name="r", servings=2, ingredients=[row], instructions=["a"]))
        self.assertEqual(self.last()["body"]["ingredients"][0]["name"], "x")
        run(self.tool.correct_cook("c1", adjustments=['{"item_id":"i1","actual_used":2,"unit":"g"}']))
        self.assertEqual(self.last()["body"]["adjustments"], [{"item_id": "i1", "actual_used": 2, "unit": "g"}])
        run(self.tool.restock("L1", substitutions=['{"for":"a","name":"b","quantity":1,"unit":"g"}'],
                              actual=['{"name":"a","quantity":2,"unit":"g"}']))
        b = self.last()["body"]
        self.assertEqual(b["substitutions"][0]["for"], "a")
        self.assertEqual(b["actual"][0]["quantity"], 2)
        run(self.tool.record_evaluation("c1", exposures=['{"subject":"leek","reaction":"refused"}']))
        self.assertEqual(self.last()["body"]["exposures"], [{"subject": "leek", "reaction": "refused"}])
        run(self.tool.propose_preferences(statements=['{"statement":"no leeks","strength":"soft"}']))
        self.assertEqual(self.last()["body"]["statements"][0]["statement"], "no leeks")

    def test_bad_element_in_other_row_params_sends_nothing(self):
        cases = [
            lambda: self.tool.save_recipe(name="r", servings=2, ingredients=["two eggs"], instructions=["a"]),
            lambda: self.tool.correct_cook("c1", adjustments=["more oil"]),
            lambda: self.tool.restock("L1", actual=["1kg"]),
            lambda: self.tool.record_evaluation("c1", exposures=["refused"]),
            lambda: self.tool.propose_preferences(statements=["no leeks"]),
        ]
        for c in cases:
            out = json.loads(run(c()))
            self.assertEqual(out["error"], "invalid_row")
        self.assertEqual(self.svc.requests, [])

    def test_restock_bought_forms(self):
        for given, want in [(None, "all"), ("all", "all"), (["milk", "eggs"], ["milk", "eggs"]),
                            ('["milk"]', ["milk"]), (" ALL ", "all")]:
            run(self.tool.restock("L1", bought=given))
            self.assertEqual(self.last()["body"]["bought"], want, repr(given))

    def test_restock_bought_default_is_all(self):
        run(self.tool.restock("L1"))
        self.assertEqual(self.last()["body"]["bought"], "all")

    def test_restock_bought_comma_string_is_invalid_and_sends_nothing(self):
        n = len(self.svc.requests)
        for bad in ("milk, eggs", "milk", "", '{"a":1}', ["ok", 3]):
            out = json.loads(run(self.tool.restock("L1", bought=bad)))
            self.assertFalse(out["ok"], repr(bad))
            self.assertEqual(out["error"], "invalid", repr(bad))
            self.assertIn("bought", out["detail"])
        self.assertEqual(len(self.svc.requests), n)

    def test_restock_polish_normalisation(self):
        for given, want in [(["all"], "all"), ([" ALL "], "all"), ([" milk "], ["milk"]),
                            ([" milk ", "eggs"], ["milk", "eggs"]), ([], [])]:
            run(self.tool.restock("L1", bought=given))
            self.assertEqual(self.last()["body"]["bought"], want, repr(given))
        run(self.tool.restock("L1", except_items=[" leeks "]))
        self.assertEqual(self.last()["body"]["except"], ["leeks"])
        run(self.tool.restock("L1", except_items='[" leeks ", "kale"]'))
        self.assertEqual(self.last()["body"]["except"], ["leeks", "kale"])

    def test_restock_polish_blank_elements_refused(self):
        n = len(self.svc.requests)
        for kw in ({"bought": ["milk", ""]}, {"bought": ["  "]}, {"except_items": [""]},
                   {"except_items": ["leeks", "   "]}):
            out = json.loads(run(self.tool.restock("L1", **kw)))
            self.assertFalse(out["ok"], repr(kw))
            self.assertEqual(out["error"], "invalid", repr(kw))
            self.assertTrue(out["nothing_sent"], repr(kw))
        self.assertEqual(len(self.svc.requests), n)

    def test_restock_docstrings(self):
        doc = pantry.Tools.restock.__doc__
        self.assertNotIn('bought="all"', doc)
        self.assertIn("empty list", doc.lower())
        self.assertIn("nothing was bought", doc.lower())

    def test_string_lists_untouched(self):
        run(self.tool.save_recipe(name="r", servings=2, ingredients=[{"name": "x"}], instructions=["a", "b"], tags=["t"]))
        b = self.last()["body"]
        self.assertEqual(b["instructions"], ["a", "b"])
        self.assertEqual(b["tags"], ["t"])


class Errors(Base):
    def test_401(self):
        self.svc.routes[("GET", "/pantry")] = (401, {"error": "unauthorized"})
        out = json.loads(run(self.tool.get_pantry()))
        self.assertFalse(out["ok"])
        self.assertEqual((out["http_status"], out["error"]), (401, "unauthorized"))
        self.assertIn("api_key", out["instruction"])

    def test_409_allergen_conflict(self):
        body = {"error": "allergen_conflict", "detail": "pesto has nuts",
                "conflicts": [{"ingredient": "pesto", "allergen": "tree nuts", "who": "Kid"}]}
        self.svc.routes[("POST", "/cook")] = (409, body)
        out = json.loads(run(self.tool.cook("r1")))
        self.assertFalse(out["ok"])
        self.assertEqual(out["error"], "allergen_conflict")
        self.assertEqual(out["conflicts"], body["conflicts"])
        self.assertIn("NOTHING was written", out["instruction"])

    def test_unmatched_and_unconvertible_surface(self):
        self.svc.routes[("POST", "/restock")] = (200, {"restocked": [], "unmatched": [{"name": "thighs", "candidates": ["chicken"]}], "unconvertible": [{"name": "x"}]})
        out = json.loads(run(self.tool.restock("l1")))
        self.assertEqual(len(out["attention"]), 2)

    def test_unreachable(self):
        self.tool.valves.service_url = "http://127.0.0.1:1"
        out = json.loads(run(self.tool.get_pantry()))
        self.assertEqual(out["error"], "service_unreachable")

    def test_guard_blocks_foreign_host(self):
        import urllib.request

        with self.assertRaises(Exception):
            urllib.request.urlopen("http://192.0.2.10:9/", timeout=1)
        self.assertTrue(BAD_CONNECTS)
        BAD_CONNECTS.clear()  # the guard did its job; tearDown asserts nothing else leaked


def make_csv(n, headers=None):
    buf = io.StringIO(newline="")
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(headers or pantry.EXPORT_COLUMNS)
    for i in range(n):
        w.writerow([f"id{i}", f"Item {i}, \"special\" ü", "Cat", "counted", i + 1, "g", "", "shelf", "2027-01-01", "milk;soy", ""])
    return buf.getvalue().encode("utf-8")


def make_xlsx(n):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(pantry.EXPORT_COLUMNS)
    for i in range(n):
        ws.append([f"id{i}", f"Item {i}", "Cat", "counted", i + 1, "g", None, "shelf", "2027-01-01", "milk;soy", None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


class Import(Base):
    def test_csv_600_rows_exact(self):
        self.preview_echo()
        files = self.attach("big.csv", make_csv(600))
        out = json.loads(run(self.tool.import_pantry(__files__=files)))
        r = self.last()
        self.assertEqual((r["method"], r["path"]), ("POST", "/audit/preview"))
        self.assertEqual(r["headers"]["x-pantry-key"], "k-test")
        self.assertEqual(len(r["body"]["rows"]), 600)
        self.assertEqual(r["body"]["source"], "csv")
        self.assertEqual(r["body"]["file_name"], "big.csv")
        self.assertEqual([x["id"] for x in r["body"]["rows"]], [f"id{i}" for i in range(600)])
        self.assertEqual(r["body"]["rows"][5]["name"], 'Item 5, "special" ü')
        self.assertEqual(r["body"]["rows"][5]["allergens"], ["milk", "soy"])
        self.assertEqual(r["body"]["rows"][5]["quantity"], 6)
        self.assertEqual(out["echo_row_count"], 600)
        self.assertEqual(out["rows_sent"], 600)
        self.assertEqual(self.svc.requests.count(r), 1)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed (run the container command)")
    def test_xlsx_600_rows_exact(self):
        self.preview_echo()
        files = self.attach("big.xlsx", make_xlsx(600))
        out = json.loads(run(self.tool.import_pantry(__files__=files)))
        r = self.last()
        self.assertEqual(r["body"]["source"], "xlsx")
        self.assertEqual(len(r["body"]["rows"]), 600)
        self.assertEqual([x["id"] for x in r["body"]["rows"]], [f"id{i}" for i in range(600)])
        self.assertEqual(r["body"]["rows"][599]["quantity"], 600)
        self.assertEqual(out["echo_row_count"], 600)

    def test_actual_column_overrides_quantity(self):
        self.preview_echo()
        data = "id,name,kind,quantity,unit,level,actual\ni1,Chickpeas,counted,4,count,,2\ni2,Oil,staple,,,low,out\ni3,Rice,counted,500,g,,\n"
        run(self.tool.import_pantry(__files__=self.attach("a.csv", data.encode())))
        rows = self.last()["body"]["rows"]
        self.assertEqual(rows[0]["quantity"], 2)
        self.assertEqual(rows[1]["level"], "out")
        self.assertEqual(rows[2]["quantity"], 500)

    def test_header_mismatch_proposes_mapping_and_posts_nothing(self):
        self.preview_echo()
        data = "Item,Qty,UOM,Shelf Life Notes\nRice,500,g,x\nBeans,2,count,y\n"
        files = self.attach("mine.csv", data.encode())
        out = json.loads(run(self.tool.import_pantry(__files__=files)))
        self.assertEqual(out["needs_confirmation"], "column_mapping")
        self.assertEqual(out["proposed_mapping"], {"Item": "name", "Qty": "quantity", "UOM": "unit", "Shelf Life Notes": None})
        self.assertEqual(out["row_count"], 2)
        self.assertEqual(self.svc.requests, [], "nothing may be posted before the mapping is confirmed")
        # confirmed mapping -> now it is sent, with the mapped names
        mp = json.dumps({"Item": "name", "Qty": "quantity", "UOM": "unit", "Shelf Life Notes": None})
        run(self.tool.import_pantry(mapping=mp, __files__=files))
        self.assertEqual(len(self.svc.requests), 1)
        self.assertEqual(self.last()["body"]["rows"], [{"name": "Rice", "quantity": 500, "unit": "g"}, {"name": "Beans", "quantity": 2, "unit": "count"}])

    def test_bad_quantity_is_not_sent(self):
        data = "name,quantity\nRice,lots\n"
        out = json.loads(run(self.tool.import_pantry(__files__=self.attach("b.csv", data.encode()))))
        self.assertFalse(out["ok"])
        self.assertEqual(self.svc.requests, [])

    def test_no_file(self):
        out = json.loads(run(self.tool.import_pantry(__files__=[])))
        self.assertEqual(out["error"], "no_file")

    def test_preview_is_capped_for_the_model(self):
        self.tool.valves.preview_max_lines = 5
        self.svc.routes[("POST", "/audit/preview")] = (200, {"preview_id": "p", "changed": [{"id": str(i)} for i in range(50)], "new": [], "missing": [], "unmatched": [], "unconvertible": []})
        out = json.loads(run(self.tool.import_pantry(__files__=self.attach("c.csv", make_csv(3)))))
        self.assertEqual(len(out["changed"]), 5)
        self.assertEqual(out["changed_count"], 50)
        self.assertTrue(out["changed_truncated"])

    def test_commit_only_via_confirm(self):
        self.preview_echo()
        run(self.tool.import_pantry(__files__=self.attach("c.csv", make_csv(3))))
        self.assertTrue(all(r["path"] != "/audit/commit" for r in self.svc.requests))

    def test_preview_expired_surfaces(self):
        self.svc.routes[("POST", "/audit/commit")] = (410, {"error": "preview_expired", "detail": "gone"})
        out = json.loads(run(self.tool.confirm_import("pv1")))
        self.assertEqual(out["error"], "preview_expired")
        self.assertIn("again", out["instruction"])


UNPLAIN = [
    "1,000", "1.000", "12.500", "1,5", "$5", "5 kg", "5kg", "nan", "NaN", "inf", "-inf", "-3", "-0.5",
    "1e3", "5.", "+5", "0x10", "1_000", "", "1 000", "1,000.50", "£1", "lots",
    "٣",  # ARABIC-INDIC DIGIT THREE
    "５",  # FULLWIDTH DIGIT FIVE
    "5\n6",
]
PLAIN = {"0": 0, "2": 2, ".25": 0.25, "0.125": 0.125, "1.5": 1.5, "1000": 1000, "12.34": 12.34, " 5 ": 5}


class NumberParsing(unittest.TestCase):
    """_to_number must never guess."""

    def test_unplain_numbers_refused(self):
        for bad in UNPLAIN:
            with self.subTest(value=bad):
                self.assertIsNone(pantry.Tools._to_number(bad))

    def test_plain_numbers_accepted(self):
        for text, want in PLAIN.items():
            with self.subTest(value=text):
                self.assertEqual(pantry.Tools._to_number(text), want)

    def test_xlsx_numbers_cells(self):
        cell = pantry.Tools._cell_text
        # numeric cells are unambiguous: 1.125 and 12.500 (a float) are real numbers
        for v, want in ((1000, 1000), (1000.0, 1000), (1.125, 1.125), (12.5, 12.5), (0, 0), (0.5, 0.5)):
            with self.subTest(cell=v):
                self.assertEqual(pantry.Tools._to_number(cell(v)), want)
        # negative, nan, inf, bool and text cells that are not plain numbers are refused
        for v in (-3, -0.5, -3.0, float("nan"), float("inf"), float("-inf"), True, False, "1,000", "1.000", "$5", "5 kg", "abc"):
            with self.subTest(cell=v):
                self.assertIsNone(pantry.Tools._to_number(cell(v)))

    def test_ambiguous_text_but_not_numeric_cell(self):
        self.assertIsNone(pantry.Tools._to_number("1.000"))
        self.assertEqual(pantry.Tools._to_number(pantry._Num("1.000")), 1.0)


class NumberImport(Base):
    def _csv(self, bad, col="quantity"):
        if col == "quantity":
            return f'name,quantity,actual\nGood,2,\nRice,"{bad}",\nBeans,1,\n'
        return f'name,quantity,actual\nGood,2,\nRice,1,"{bad}"\nBeans,1,\n'

    def test_unplain_numbers_never_posted_end_to_end(self):
        self.preview_echo()
        for bad in UNPLAIN:
            if bad in ("", "5\n6"):  # an empty cell is simply "no value"; a newline needs quoting tricks
                continue
            for col in ("quantity", "actual"):
                with self.subTest(value=bad, column=col):
                    self.svc.requests.clear()
                    out = json.loads(run(self.tool.import_pantry(__files__=self.attach("n.csv", self._csv(bad, col).encode("utf-8")))))
                    self.assertFalse(out["ok"])
                    self.assertTrue(out["nothing_sent"])
                    self.assertEqual(self.svc.requests, [])
                    self.assertEqual([(i["row"], i["column"], i["value"]) for i in out["invalid"]], [(3, col, bad)])

    def test_plain_numbers_accepted_end_to_end(self):
        self.preview_echo()
        data = "name,quantity\n" + "".join(f"n{i},{t}\n" for i, t in enumerate(k for k in PLAIN if k.strip() == k))
        run(self.tool.import_pantry(__files__=self.attach("ok.csv", data.encode())))
        want = [v for k, v in PLAIN.items() if k.strip() == k]
        self.assertEqual([r["quantity"] for r in self.last()["body"]["rows"]], want)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed (run the container command)")
    def test_xlsx_numbers_end_to_end(self):
        def sheet(vals):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.append(["name", "quantity"])
            for i, v in enumerate(vals):
                ws.append([f"n{i}", v])
            out = io.BytesIO()
            wb.save(out)
            return out.getvalue()

        self.preview_echo()
        run(self.tool.import_pantry(__files__=self.attach("g.xlsx", sheet([1000, 1.125, 0.5, 3.0, 0]))))
        self.assertEqual([r["quantity"] for r in self.last()["body"]["rows"]], [1000, 1.125, 0.5, 3, 0])
        # (openpyxl cannot store nan/inf in a cell; those are covered at cell level in test_xlsx_numbers_cells)
        for bad in (-3, -0.5, True, "1,000", "1.000", "$5", "5 kg", "nan", "٣"):
            with self.subTest(value=bad):
                self.svc.requests.clear()
                out = json.loads(run(self.tool.import_pantry(__files__=self.attach("b.xlsx", sheet([2, bad])))))
                self.assertFalse(out["ok"])
                self.assertTrue(out["nothing_sent"])
                self.assertEqual(self.svc.requests, [])
                self.assertEqual(out["invalid"][0]["row"], 3)


class StapleLevels(Base):
    def _post(self, data):
        self.svc.requests.clear()
        return json.loads(run(self.tool.import_pantry(__files__=self.attach("s.csv", data.encode()))))

    def test_staple_actual_must_be_a_level(self):
        self.preview_echo()
        for word, want in (("LOW", "low"), (" Plenty ", "plenty"), ("out", "out")):
            with self.subTest(actual=word):
                out = self._post(f'name,kind,level,actual\nOil,staple,plenty,"{word}"\n')
                self.assertTrue(out["ok"])
                self.assertEqual(self.last()["body"]["rows"], [{"name": "Oil", "kind": "staple", "level": want}])
        for bad in ("5", "lowish", "medium", "1,000", "-1", "nan"):
            with self.subTest(actual=bad):
                out = self._post(f'name,kind,level,actual\nOil,staple,plenty,"{bad}"\n')
                self.assertFalse(out["ok"])
                self.assertTrue(out["nothing_sent"])
                self.assertEqual(self.svc.requests, [])
                self.assertEqual([(i["row"], i["column"], i["value"]) for i in out["invalid"]], [(2, "actual", bad)])

    def test_counted_actual_level_word_is_invalid_not_posted_as_both(self):
        for word in ("LOW", "out", "plenty"):
            with self.subTest(actual=word):
                out = self._post(f"name,kind,quantity,unit,actual\nRice,counted,500,g,{word}\n")
                self.assertFalse(out["ok"])
                self.assertEqual(self.svc.requests, [])
                self.assertEqual([(i["row"], i["column"], i["value"]) for i in out["invalid"]], [(2, "actual", word)])

    def test_level_column_must_be_a_level(self):
        self.preview_echo()
        for bad in ("Medium", "half", "5", "lo"):
            with self.subTest(level=bad):
                out = self._post(f"name,kind,level\nOil,staple,{bad}\n")
                self.assertFalse(out["ok"])
                self.assertEqual(self.svc.requests, [])
                self.assertEqual([(i["row"], i["column"], i["value"]) for i in out["invalid"]], [(2, "level", bad)])
        self._post("name,kind,level\nOil,staple, LOW \n")
        self.assertEqual(self.last()["body"]["rows"][0]["level"], "low")

    def test_unknown_kind_level_word_is_a_level(self):
        self.preview_echo()
        self._post("id,name,actual\ni2,Oil,out\n")
        self.assertEqual(self.last()["body"]["rows"], [{"id": "i2", "name": "Oil", "level": "out"}])


class SkillMatchesTool(unittest.TestCase):
    def test_write_rule_lists_exactly_the_write_functions(self):
        import inspect
        import re

        with open(os.path.join(HERE, "..", "..", "skills", "kitchen-pantry.md"), encoding="utf-8") as fh:
            rule = next(l for l in fh.read().splitlines() if l.startswith("4. **Show every write back.**"))
        listed = set(re.findall(r"`([a-z_]+)`", rule.split("(", 1)[1].split(")", 1)[0]))
        read_only_posts = {"get_guidance", "import_pantry"}  # POST routes that write nothing
        writers = set()
        for name, fn in inspect.getmembers(pantry.Tools, inspect.iscoroutinefunction):
            if name.startswith("_") or name in read_only_posts:
                continue
            src = inspect.getsource(fn)
            if '"POST"' in src or '"PUT"' in src or "self._save_recipe(" in src:
                writers.add(name)
        self.assertEqual(listed, writers)
        self.assertTrue({"log_cooked_meal", "correct_cook", "cook", "set_plan_status", "confirm_import"} <= listed)


class Export(Base):
    def setUp(self):
        super().setUp()
        self.svc.routes[("GET", "/pantry")] = (200, {"items": STOCK})
        self.delivered = []

        async def fake_deliver(name, data, ctype, user, emitter):
            self.delivered.append((name, data, ctype))
            return {"id": "fid", "url": "/api/v1/files/fid/content"}

        self.tool._deliver_file = fake_deliver

    def test_default_is_csv_with_columns(self):
        msg = run(self.tool.export_pantry())
        name, data, ctype = self.delivered[0]
        self.assertTrue(name.endswith(".csv"))
        self.assertEqual(ctype, "text/csv")
        self.assertIn("/api/v1/files/fid/content", msg)
        rows = list(csv.reader(io.StringIO(data.decode("utf-8-sig"))))
        self.assertEqual(rows[0], ["id", "name", "category", "kind", "quantity", "unit", "level", "location", "expires_on", "allergens", "actual"])
        self.assertEqual(len(rows), 4)
        self.assertEqual(rows[3][9], "tree nuts;milk")
        self.assertTrue(all(r[10] == "" for r in rows[1:]))

    @unittest.skipIf(openpyxl is None, "openpyxl not installed (run the container command)")
    def test_xlsx_option(self):
        run(self.tool.export_pantry(format="xlsx"))
        name, data, ctype = self.delivered[0]
        self.assertTrue(name.endswith(".xlsx"))
        ws = openpyxl.load_workbook(io.BytesIO(data)).active
        self.assertEqual([c.value for c in ws[1]][:3], ["id", "name", "category"])
        self.assertEqual(ws.max_row, 4)

    def test_fallback_inline_fenced_csv(self):
        async def none(*a, **k):
            return None

        self.tool._deliver_file = none
        msg = run(self.tool.export_pantry())
        self.assertIn("```csv\nid,name,category", msg)
        self.assertIn("Pesto", msg)

    def test_deliver_file_without_owui_returns_none(self):
        from unittest import mock

        blocked = {"open_webui": None, "open_webui.models": None, "open_webui.models.files": None,
                   "open_webui.storage": None, "open_webui.storage.provider": None}
        with mock.patch.dict(sys.modules, blocked):
            t = pantry.Tools()
            self.assertIsNone(run(t._deliver_file("a.csv", b"x", "text/csv", {"id": "u"}, None)))

    def test_deliver_file_with_stubbed_owui_emits_files_event(self):
        import types
        from unittest import mock

        stored = {}
        events = []

        class FileForm:
            def __init__(self, **kw):
                self.kw = kw

        class Files:
            @staticmethod
            async def insert_new_file(user_id, form):
                stored["row"] = (user_id, form.kw)
                return object()

        class Storage:
            @staticmethod
            def upload_file(fh, filename, tags):
                data = fh.read()
                stored["bytes"] = data
                return data, "/data/uploads/" + filename

        mods = {
            "open_webui": types.ModuleType("open_webui"),
            "open_webui.models": types.ModuleType("open_webui.models"),
            "open_webui.models.files": types.SimpleNamespace(Files=Files, FileForm=FileForm),
            "open_webui.storage": types.ModuleType("open_webui.storage"),
            "open_webui.storage.provider": types.SimpleNamespace(Storage=Storage),
        }

        async def emitter(ev):
            events.append(ev)

        with mock.patch.dict(sys.modules, mods):
            t = pantry.Tools()
            got = run(t._deliver_file("p.csv", b"abc", "text/csv", {"id": "u1"}, emitter))
        self.assertEqual(stored["bytes"], b"abc")
        self.assertEqual(stored["row"][0], "u1")
        self.assertEqual(stored["row"][1]["meta"]["name"], "p.csv")
        self.assertEqual(got["url"], f"/api/v1/files/{got['id']}/content")
        self.assertEqual(events, [{"type": "files", "data": {"files": [{"type": "file", "url": got["url"], "name": "p.csv"}]}}])

    def test_bad_format(self):
        out = json.loads(run(self.tool.export_pantry(format="pdf")))
        self.assertEqual(out["error"], "invalid")

    def _expected_rows(self, fmt_rows):
        exp = []
        for it in STOCK:
            r = {"id": it["id"], "name": it["name"], "kind": it["kind"]}
            for k in ("category", "unit", "level", "location", "expires_on"):
                if it.get(k):
                    r[k] = it[k]
            if it.get("quantity") is not None:
                q = it["quantity"]
                r["quantity"] = int(q) if float(q).is_integer() else q
            if it["allergens"]:
                r["allergens"] = it["allergens"]
            exp.append(r)
        return exp

    def test_round_trip_csv(self):
        run(self.tool.export_pantry())
        data = self.delivered[0][1]
        self.preview_echo()
        out = json.loads(run(self.tool.import_pantry(__files__=self.attach("pantry.csv", data))))
        self.assertTrue(out["ok"])
        self.assertEqual(self.last()["body"]["rows"], self._expected_rows(None))

    @unittest.skipIf(openpyxl is None, "openpyxl not installed (run the container command)")
    def test_round_trip_xlsx(self):
        run(self.tool.export_pantry(format="xlsx"))
        data = self.delivered[0][1]
        self.preview_echo()
        run(self.tool.import_pantry(__files__=self.attach("pantry.xlsx", data)))
        self.assertEqual(self.last()["body"]["rows"], self._expected_rows(None))


if __name__ == "__main__":
    unittest.main()
