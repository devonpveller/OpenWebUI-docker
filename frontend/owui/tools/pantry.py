"""
title: Pantry (Kitchen model)
author: ai-stack / Open Brain
version: 1.0.0
description: >
  Thin OWUI client for the openbrain-pantry service: pantry stock, recipes, cook /
  plan / shopping / restock flows, evaluations and taste preferences, household
  members, and CSV / xlsx exchange of the pantry. ALL arithmetic and rules live in
  the service; every function here is one documented HTTP call. The only logic in
  this file is spreadsheet parsing/building (CSV default, xlsx via openpyxl) and
  turning service errors into messages the model must relay instead of guessing.
  Pair it with the `kitchen-pantry` skill (the Kitchen model's instructions).
requirements: openpyxl
"""

import asyncio
import csv
import datetime
import io
import json
import re
import urllib.error
import urllib.request
import uuid
from typing import Any, Awaitable, Callable, Optional

from pydantic import BaseModel, Field

# Columns of the exported sheet (PLAN flow 8). `actual` is the empty column the
# household fills in during a thorough audit.
EXPORT_COLUMNS = [
    "id",
    "name",
    "category",
    "kind",
    "quantity",
    "unit",
    "level",
    "location",
    "expires_on",
    "allergens",
    "actual",
]
_ROW_FIELDS = [c for c in EXPORT_COLUMNS if c != "actual"]
_NUMERIC_FIELDS = {"quantity"}
_LIST_FIELDS = {"allergens", "may_contain"}
# Extra, optional sheet columns that map straight onto an /audit/preview row key.
_OPTIONAL_COLUMNS = ["may_contain"]
_CANONICAL = EXPORT_COLUMNS + _OPTIONAL_COLUMNS

# Suggestions only - a header that is not exactly canonical is NEVER applied
# without the user confirming the mapping.
_SYNONYMS = {
    "item": "name",
    "product": "name",
    "ingredient": "name",
    "item name": "name",
    "qty": "quantity",
    "amount": "quantity",
    "count": "quantity",
    "on hand": "quantity",
    "stock": "quantity",
    "uom": "unit",
    "units": "unit",
    "type": "kind",
    "section": "category",
    "aisle": "category",
    "place": "location",
    "where": "location",
    "storage": "location",
    "expiry": "expires_on",
    "expires": "expires_on",
    "expiry date": "expires_on",
    "best before": "expires_on",
    "use by": "expires_on",
    "allergen": "allergens",
    "allergy": "allergens",
    "counted": "actual",
    "actual count": "actual",
    "real": "actual",
}


class _Num(str):
    """Text of a numeric spreadsheet cell (xlsx int/float): unambiguous, unlike csv text."""


_PLAIN_NUMBER = re.compile(r"(?:[0-9]+(?:\.[0-9]+)?|\.[0-9]+)")  # ASCII digits only; used with fullmatch
_LEVELS = ("plenty", "low", "out")
# "1.000" / "12.500": could be a thousands separator in some locales - ask, do not guess.
_AMBIGUOUS_THOUSANDS = re.compile(r"[1-9][0-9]{0,2}\.[0-9]{3}")


class ServiceError(Exception):
    def __init__(self, status: int, body: Any):
        super().__init__(f"HTTP {status}")
        self.status = status
        self.body = body


def _compact(obj: Any) -> str:
    return json.dumps(obj, ensure_ascii=False, separators=(",", ":"), default=str)


def _coerce(value: Any) -> Any:
    """Local models sometimes send list/dict arguments as JSON strings."""
    if isinstance(value, str):
        s = value.strip()
        if s[:1] in "[{":
            try:
                return json.loads(s)
            except ValueError:
                return value
    return value


class InvalidRow(Exception):
    """A string element of a row-list parameter that is not one JSON object."""

    def __init__(self, param: str, index: int, value: Any):
        super().__init__(f"{param}[{index}]")
        self.param = param
        self.index = index
        self.value = value

    def result(self) -> dict:
        shown = self.value if isinstance(self.value, str) else _compact(self.value)
        if len(shown) > 200:
            shown = shown[:200] + "..."
        return {
            "ok": False,
            "error": "invalid_row",
            "nothing_sent": True,
            "detail": f"{self.param}[{self.index}] is not a JSON object: {shown!r}",
            "instruction": (
                "Nothing was sent. Every row must be an object like {\"name\": ..., ...}. "
                "Rebuild that row from what the user said (ask if unclear) and call again; do not guess."
            ),
        }


def _rows(value: Any, param: str) -> Any:
    """A list-of-objects argument. Accepts the whole list as a JSON string, and string
    ELEMENTS that each hold one JSON object (local models send those). An element that
    is not an object raises InvalidRow - nothing is guessed and nothing is sent."""
    value = _coerce(value)
    if not isinstance(value, list):
        return value
    out = []
    for i, el in enumerate(value):
        if isinstance(el, str):
            try:
                parsed = json.loads(el.strip())
            except ValueError:
                raise InvalidRow(param, i, el)
            if not isinstance(parsed, dict):
                raise InvalidRow(param, i, el)
            el = parsed
        elif not isinstance(el, dict):
            raise InvalidRow(param, i, el)
        out.append(el)
    return out


def _clean(d: dict) -> dict:
    return {k: v for k, v in d.items() if v is not None}


class Tools:
    class Valves(BaseModel):
        service_url: str = Field(
            default="http://openbrain-pantry:8000",
            description="Base URL of the openbrain-pantry service.",
        )
        api_key: str = Field(
            default="",
            description="PANTRY_API_KEY - sent as the x-pantry-key header on every call.",
        )
        request_timeout_s: float = Field(
            default=60.0, description="Per-request timeout in seconds."
        )
        list_limit: int = Field(
            default=200,
            description="Max items returned to the model by get_pantry (the rest is reported as truncated).",
        )
        preview_max_lines: int = Field(
            default=40,
            description="Max lines per list (changed/new/missing/...) shown to the model from an import preview.",
        )

    def __init__(self):
        self.valves = self.Valves()

    # ------------------------------------------------------------------ http

    def _request_sync(
        self, method: str, path: str, body: Any = None, query: Optional[dict] = None
    ) -> Any:
        v = self.valves
        url = v.service_url.rstrip("/") + path
        if query:
            q = {k: val for k, val in query.items() if val not in (None, "")}
            if q:
                from urllib.parse import urlencode

                url += "?" + urlencode(
                    {k: str(x).lower() if isinstance(x, bool) else x for k, x in q.items()}
                )
        data = None
        headers = {"x-pantry-key": v.api_key, "Accept": "application/json"}
        if body is not None:
            data = json.dumps(body, ensure_ascii=False).encode("utf-8")
            headers["Content-Type"] = "application/json"
        req = urllib.request.Request(url, data=data, method=method, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=v.request_timeout_s) as resp:
                raw = resp.read()
                status = resp.status
        except urllib.error.HTTPError as e:
            raw = e.read()
            status = e.code
            try:
                parsed = json.loads(raw.decode("utf-8")) if raw else {}
            except ValueError:
                parsed = {"error": "http_error", "detail": raw[:300].decode("utf-8", "replace")}
            raise ServiceError(status, parsed)
        try:
            return json.loads(raw.decode("utf-8")) if raw else {}
        except ValueError:
            raise ServiceError(status, {"error": "bad_response", "detail": "service returned non-JSON"})

    async def _call(
        self, method: str, path: str, body: Any = None, query: Optional[dict] = None
    ) -> Any:
        return await asyncio.to_thread(self._request_sync, method, path, body, query)

    @staticmethod
    def _error_result(e: Exception) -> dict:
        if isinstance(e, ServiceError):
            body = e.body if isinstance(e.body, dict) else {"detail": str(e.body)}
            code = body.get("error", "http_error")
            out = {"ok": False, "http_status": e.status, **body}
            out["error"] = code
            if e.status == 401:
                out["instruction"] = (
                    "The pantry key is missing or wrong. Tell the user to set the api_key valve "
                    "of the Pantry tool. Do not retry and do not guess."
                )
            elif code == "allergen_conflict":
                out["instruction"] = (
                    "NOTHING was written. Tell the user which ingredient conflicts with whose "
                    "allergy. Never retry with that ingredient; propose a substitute and ask."
                )
            elif code == "preview_expired":
                out["instruction"] = "The preview expired. Run import_pantry again and re-confirm."
            elif e.status == 409:
                out["instruction"] = "The service refused this as a conflict. Tell the user; do not retry blindly."
            elif e.status >= 500:
                out["instruction"] = "The pantry service failed. Tell the user; do not invent a result."
            else:
                out["instruction"] = "The service rejected the request. Read detail, fix the input or ask the user."
            return out
        return {
            "ok": False,
            "error": "service_unreachable",
            "detail": f"{type(e).__name__}: {e}",
            "instruction": "The pantry service could not be reached. Tell the user; do not invent a result.",
        }

    @staticmethod
    def _attention(body: Any) -> Any:
        """Name the lists the model must ask the user about, never resolve itself."""
        if not isinstance(body, dict):
            return body
        notes = []
        if body.get("unmatched"):
            notes.append(
                "unmatched: ask the user which pantry item each one is (use candidates); never pick one yourself"
            )
        if body.get("unconvertible"):
            notes.append(
                "unconvertible: those lines were NOT written (unit dimension mismatch); ask the user for a usable quantity"
            )
        if body.get("allergen_conflicts"):
            notes.append("allergen_conflicts: tell the user; remove or substitute before the recipe is used")
        if body.get("shortfalls"):
            notes.append("shortfalls: tell the user what is short; offer the shopping list")
        if notes:
            body = {**body, "attention": notes}
        return body

    async def _do(
        self,
        method: str,
        path: str,
        body: Any = None,
        query: Optional[dict] = None,
        emitter: Optional[Callable[[dict], Awaitable[None]]] = None,
        status: str = "",
    ) -> str:
        await self._emit(emitter, status or f"{method} {path}", False)
        try:
            res = await self._call(method, path, body, query)
        except Exception as e:  # ServiceError, URLError, timeout
            await self._emit(emitter, "Pantry service error", True)
            return _compact(self._error_result(e))
        await self._emit(emitter, "Done", True)
        return _compact(self._attention(res))

    @staticmethod
    async def _emit(emitter, desc: str, done: bool):
        if emitter:
            try:
                await emitter({"type": "status", "data": {"description": desc, "done": done}})
            except Exception:
                pass

    # ------------------------------------------------------------- pantry

    async def get_pantry(
        self,
        q: Optional[str] = None,
        category: Optional[str] = None,
        use_soon: Optional[bool] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Read the pantry: items with quantity, unit, level, reserved and AVAILABLE amounts, expiry and allergens.
        Always take recipe ingredient ids from here. Use the `available` field, not `quantity`, when planning.

        :param q: Optional text filter on item name.
        :param category: Optional category filter.
        :param use_soon: True to list only items that should be used soon.
        :return: JSON with the items.
        """
        try:
            res = await self._call(
                "GET", "/pantry", None, {"q": q, "category": category, "use_soon": use_soon}
            )
        except Exception as e:
            return _compact(self._error_result(e))
        items = res.get("items", []) if isinstance(res, dict) else []
        lim = max(1, int(self.valves.list_limit))
        if len(items) > lim:
            return _compact(
                {
                    "items": items[:lim],
                    "total": len(items),
                    "truncated": True,
                    "note": "Too many items to show. Narrow with q or category.",
                }
            )
        return _compact({"items": items, "total": len(items)})

    async def update_pantry(
        self,
        items: list[dict],
        reason: str = "manual",
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Change pantry stock from what the user told you. Show the user the table of changes and get a yes BEFORE calling.
        A name that is not in the pantry is created only when that line has create=true; otherwise it comes back as unmatched with candidates - ask the user.

        :param items: List of lines, each {id or name, quantity (absolute) or delta (+/-), unit, kind ("counted"/"staple"), level ("plenty"/"low"/"out" for staples), category, location, expires_on, allergens, may_contain, aliases, create (true only for a new item)}.
        :param reason: "manual" for stock entry, "correct" to fix a wrong number.
        :return: JSON with applied, created, unmatched and unconvertible lines.
        """
        try:
            rows = _rows(items, "items")
        except InvalidRow as e:
            return _compact(e.result())
        return await self._do(
            "POST",
            "/pantry/adjust",
            {"reason": reason, "items": rows},
            emitter=__event_emitter__,
            status="Updating pantry",
        )

    # ----------------------------------------------------------- guidance

    async def get_guidance(
        self,
        theme: Optional[str] = None,
        recipe_id: Optional[str] = None,
        guest_context: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Fetch taste and safety guidance. Call it BEFORE proposing themes, writing or revising a recipe, or asking a curiosity question.
        Returns allergens to exclude (household + guests), hard exclusions, contextual dislikes, soft signals, child cooldowns and trends, hypotheses, recent evaluations, use-soon items and untried cuisines/techniques/ingredients.

        :param theme: The theme being considered, if any.
        :param recipe_id: The recipe being revised or evaluated, if any.
        :param guest_context: Guests for this ONE meal: {adults, children, allergies, avoid, diet, note}.
        :return: JSON guidance.
        """
        return await self._do(
            "POST",
            "/guidance",
            _clean({"theme": theme, "recipe_id": recipe_id, "guest_context": _coerce(guest_context)}),
            emitter=__event_emitter__,
            status="Reading guidance",
        )

    # ------------------------------------------------------------ recipes

    async def _save_recipe(self, source: str, rotation, **kw) -> str:
        emitter = kw.pop("emitter", None)
        try:
            ingredients = _rows(kw.get("ingredients"), "ingredients")
        except InvalidRow as e:
            return _compact(e.result())
        body = _clean(
            {
                "id": kw.get("recipe_id"),
                "name": kw.get("name"),
                "theme": kw.get("theme"),
                "cuisine": kw.get("cuisine"),
                "servings": kw.get("servings"),
                "ingredients": ingredients,
                "instructions": _coerce(kw.get("instructions")),
                "tags": _coerce(kw.get("tags")),
                "source": source,
                "rotation": rotation,
                "draft": kw.get("draft"),
                "reason": kw.get("reason"),
                "guest_context": _coerce(kw.get("guest_context")),
            }
        )
        return await self._do("POST", "/recipes", body, emitter=emitter, status="Saving recipe")

    async def save_recipe(
        self,
        name: str,
        servings: float,
        ingredients: list[dict],
        instructions: list,
        theme: Optional[str] = None,
        cuisine: Optional[str] = None,
        tags: Optional[list] = None,
        draft: Optional[bool] = None,
        recipe_id: Optional[str] = None,
        reason: Optional[str] = None,
        guest_context: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Save a recipe YOU generated (draft or final). With recipe_id it writes a NEW revision; a cooked revision is never edited.
        Ingredients must reference pantry ids from get_pantry; any `unmatched` or `allergen_conflicts` in the result must be shown to the user.

        :param name: Recipe name.
        :param servings: Portions the recipe makes.
        :param ingredients: List of {pantry_item_id, name, quantity, unit, staple}.
        :param instructions: List of step strings.
        :param theme: Theme it belongs to.
        :param cuisine: Cuisine, for the exploration map.
        :param tags: Tags; prefix "technique:" for a technique.
        :param draft: True while the user is still iterating.
        :param recipe_id: Existing recipe id to write a new revision of.
        :param reason: Why this revision (for example the restock change).
        :param guest_context: Guests for this one meal only.
        :return: JSON with recipe, revision, unmatched and allergen_conflicts.
        """
        return await self._save_recipe(
            "generated", None, name=name, servings=servings, ingredients=ingredients,
            instructions=instructions, theme=theme, cuisine=cuisine, tags=tags, draft=draft,
            recipe_id=recipe_id, reason=reason, guest_context=guest_context,
            emitter=__event_emitter__,
        )

    async def save_household_recipe(
        self,
        name: str,
        servings: float,
        ingredients: list[dict],
        instructions: list,
        rotation: bool = False,
        cuisine: Optional[str] = None,
        tags: Optional[list] = None,
        recipe_id: Optional[str] = None,
        reason: Optional[str] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Save a recipe the HOUSEHOLD brought (pasted text or an ingredient list) so it can be reused by name. Map each ingredient to a pantry id first; ask the user about anything unmatched.

        :param name: Recipe name as the household calls it.
        :param servings: Portions it makes.
        :param ingredients: List of {pantry_item_id, name, quantity, unit, staple}.
        :param instructions: List of step strings (may be short).
        :param rotation: True if it is one of the dishes the household rotates.
        :param cuisine: Cuisine, if known.
        :param tags: Tags.
        :param recipe_id: Existing household recipe id to write a new revision of.
        :param reason: Why this revision.
        :return: JSON with recipe, revision, unmatched and allergen_conflicts.
        """
        return await self._save_recipe(
            "household", rotation, name=name, servings=servings, ingredients=ingredients,
            instructions=instructions, cuisine=cuisine, tags=tags, recipe_id=recipe_id,
            reason=reason, emitter=__event_emitter__,
        )

    # --------------------------------------------------------------- cook

    async def cook(
        self,
        recipe_id: str,
        servings: Optional[float] = None,
        meal_plan_id: Optional[str] = None,
        guest_context: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Commit tonight's dinner: deducts every ingredient in one transaction. Call ONLY after the user said "yes, we're making this one" (or a clear variant).
        Returns deductions and shortfalls (shortfalls do not block). A 409 allergen_conflict means nothing was written - tell the user.

        :param recipe_id: Recipe to cook.
        :param servings: Portions to cook; omit for the household default plus guests.
        :param meal_plan_id: The plan row being cooked, if any.
        :param guest_context: Guests for this one meal: {adults, children, allergies, avoid, diet, note}.
        :return: JSON with cook_event_id, deductions, shortfalls, unconvertible.
        """
        return await self._do(
            "POST",
            "/cook",
            _clean(
                {
                    "recipe_id": recipe_id,
                    "servings": servings,
                    "meal_plan_id": meal_plan_id,
                    "guest_context": _coerce(guest_context),
                }
            ),
            emitter=__event_emitter__,
            status="Cooking (deducting stock)",
        )

    async def log_cooked_meal(
        self,
        recipe_id: str,
        cooked_at: Optional[str] = None,
        servings: Optional[float] = None,
        guest_context: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Log a meal that was ALREADY cooked without the model ("we made tacos last night"). Deducts stock and creates a cook event that can be evaluated.

        :param recipe_id: The recipe that was cooked (create it first if needed).
        :param cooked_at: ISO date or datetime it was cooked.
        :param servings: Portions cooked, if known.
        :param guest_context: Guests at that meal, if any.
        :return: JSON like cook.
        """
        return await self._do(
            "POST",
            "/cook",
            _clean(
                {
                    "recipe_id": recipe_id,
                    "servings": servings,
                    "guest_context": _coerce(guest_context),
                    "logged_after": True,
                    "cooked_at": cooked_at,
                }
            ),
            emitter=__event_emitter__,
            status="Logging cooked meal",
        )

    async def correct_cook(
        self,
        cook_event_id: str,
        adjustments: Optional[list[dict]] = None,
        undo: bool = False,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Correct the quantities of a cook ("we used more oil") or undo it entirely.

        :param cook_event_id: The cook event to correct.
        :param adjustments: List of {item_id, actual_used, unit}.
        :param undo: True to reverse the whole cook exactly.
        :return: JSON with the corrected deductions.
        """
        try:
            rows = None if undo else _rows(adjustments, "adjustments")
        except InvalidRow as e:
            return _compact(e.result())
        body = {"undo": True} if undo else {"adjustments": rows or []}
        return await self._do(
            "POST", f"/cook/{cook_event_id}/correct", body,
            emitter=__event_emitter__, status="Correcting cook",
        )

    # --------------------------------------------------------------- plan

    async def plan_meal(
        self,
        date: str,
        recipe_id: Optional[str] = None,
        custom_meal: Optional[str] = None,
        leftovers_of: Optional[str] = None,
        servings: Optional[float] = None,
        guest_context: Optional[dict] = None,
        notes: Optional[str] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Plan one DINNER on a date. A planned recipe reserves its stock. Use custom_meal for eating out or a gap, leftovers_of for leftovers.

        :param date: YYYY-MM-DD.
        :param recipe_id: Recipe to plan.
        :param custom_meal: Free text meal when there is no recipe.
        :param leftovers_of: Plan id this meal is leftovers of.
        :param servings: Portions; omit for the household default plus guests.
        :param guest_context: Guests for this one meal.
        :param notes: Free text note.
        :return: JSON with the plan row, reservations and shortfalls.
        """
        return await self._do(
            "POST",
            "/plan",
            _clean(
                {
                    "date": date,
                    "recipe_id": recipe_id,
                    "custom_meal": custom_meal,
                    "leftovers_of": leftovers_of,
                    "servings": servings,
                    "guest_context": _coerce(guest_context),
                    "notes": notes,
                }
            ),
            emitter=__event_emitter__,
            status="Planning meal",
        )

    async def get_plan(
        self,
        date: Optional[str] = None,
        week_start: Optional[str] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Read planned dinners for one date or a week, each re-checked against current stock (shortfalls_now).

        :param date: YYYY-MM-DD for one day.
        :param week_start: YYYY-MM-DD of the week's first day for a whole week.
        :return: JSON with plans.
        """
        try:
            res = await self._call("GET", "/plan", None, {"date": date, "week_start": week_start})
        except Exception as e:
            return _compact(self._error_result(e))
        return _compact(res)

    async def set_plan_status(
        self,
        plan_id: str,
        status: str,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Mark a planned dinner "skipped" (it will not be cooked; this releases its reserved stock) or back to "planned". "cooked" is set only by cook.

        :param plan_id: The plan row id from get_plan / plan_meal.
        :param status: "skipped" or "planned".
        :return: JSON with the updated plan row.
        """
        return await self._do(
            "POST", f"/plan/{plan_id}/status", {"status": status},
            emitter=__event_emitter__, status="Updating plan status",
        )

    async def shopping_list(
        self,
        week_start: str,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Build the shopping list for a week: what planned meals need minus what is available, plus staples that are low or out, plus items carried over.

        :param week_start: YYYY-MM-DD of the week's first day.
        :return: JSON with list_id, items (name, quantity, unit, reason) and carried_over.
        """
        return await self._do(
            "POST", "/shopping-list", {"week_start": week_start},
            emitter=__event_emitter__, status="Building shopping list",
        )

    async def restock(
        self,
        list_id: str,
        bought: Optional[list[str]] = None,
        except_items: Optional[list] = None,
        substitutions: Optional[list[dict]] = None,
        actual: Optional[list[dict]] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Record what was ACTUALLY bought for a shopping list. Unbought items carry over. The result lists planned meals whose shortfalls changed; propose a revision for each.

        :param list_id: The shopping list id.
        :param bought: Item names/ids that were bought. Omit it (or pass null) when EVERYTHING on the list was bought. A list of just "all" also means everything. An empty list [] means nothing was bought.
        :param except_items: Names/ids of list items NOT bought, when everything else was bought: omit bought (or pass null) and list only the exceptions here.
        :param substitutions: List of {for, name, quantity, unit} for swapped items.
        :param actual: List of {name or id, quantity, unit} when the pack size differs from the list.
        :return: JSON with restocked, carried_over and affected_plans.
        """
        b = _coerce(bought)
        if b is None or (isinstance(b, str) and b.strip().lower() == "all"):
            b = "all"
        elif isinstance(b, list) and all(isinstance(x, str) and x.strip() for x in b):
            b = [x.strip() for x in b]
            if len(b) == 1 and b[0].lower() == "all":
                b = "all"
        else:
            return _compact(
                {"ok": False, "error": "invalid", "nothing_sent": True,
                 "detail": f"bought must be omitted (everything) or a list of non-empty item names/ids, got {_compact(bought)}",
                 "instruction": "Nothing was sent. Pass bought as a list of names, or omit it if everything was bought."}
            )
        exc = _coerce(except_items)
        if isinstance(exc, list):
            if not all(not isinstance(x, str) or x.strip() for x in exc):
                return _compact(
                    {"ok": False, "error": "invalid", "nothing_sent": True,
                     "detail": f"except_items must not contain empty names, got {_compact(except_items)}",
                     "instruction": "Nothing was sent. Pass except_items as a list of item names/ids, or omit it."}
                )
            exc = [x.strip() if isinstance(x, str) else x for x in exc]
        try:
            subs = _rows(substitutions, "substitutions")
            act = _rows(actual, "actual")
        except InvalidRow as e:
            return _compact(e.result())
        return await self._do(
            "POST",
            "/restock",
            _clean(
                {
                    "list_id": list_id,
                    "bought": b,
                    "except": exc,
                    "substitutions": subs,
                    "actual": act,
                }
            ),
            emitter=__event_emitter__,
            status="Recording restock",
        )

    # ---------------------------------------------------------- taste

    async def record_evaluation(
        self,
        cook_event_id: str,
        who: str = "all",
        rating: Optional[int] = None,
        liked: Optional[bool] = None,
        why: Optional[str] = None,
        change: Optional[str] = None,
        curiosity_q: Optional[str] = None,
        curiosity_a: Optional[str] = None,
        exposures: Optional[list[dict]] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Record how a cooked meal went. A guest meal still takes the FAMILY's evaluation; nothing guest-specific is learned.

        :param cook_event_id: The cook event being evaluated.
        :param who: "adult", "child" or "all".
        :param rating: 1-5.
        :param liked: True/False.
        :param why: The user's reason, in their words.
        :param change: What they would change.
        :param curiosity_q: The curiosity question you asked.
        :param curiosity_a: The user's answer.
        :param exposures: List of {person_id, subject, reaction ("refused"/"tolerated"/"liked")} - used for the child.
        :return: JSON with the evaluation and exposures.
        """
        try:
            expo = _rows(exposures, "exposures")
        except InvalidRow as e:
            return _compact(e.result())
        return await self._do(
            "POST",
            "/evaluations",
            _clean(
                {
                    "cook_event_id": cook_event_id,
                    "who": who,
                    "rating": rating,
                    "liked": liked,
                    "why": why,
                    "change": change,
                    "curiosity_q": curiosity_q,
                    "curiosity_a": curiosity_a,
                    "exposures": expo,
                }
            ),
            emitter=__event_emitter__,
            status="Recording evaluation",
        )

    async def propose_preferences(
        self,
        statements: list[dict],
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Store preferences you derived as PROPOSED (unconfirmed, ignored by guidance). Read each back to the user with its strength before calling confirm_preferences.

        :param statements: List of {statement, strength ("hard"/"contextual"/"soft"), subject, context, reason, scope ("recipe"/"theme"/"always"), who ("adult"/"child"/"all"), evidence}.
        :return: JSON with the proposed rows and their ids.
        """
        try:
            rows = _rows(statements, "statements")
        except InvalidRow as e:
            return _compact(e.result())
        return await self._do(
            "POST", "/preferences", {"statements": rows},
            emitter=__event_emitter__, status="Proposing preferences",
        )

    async def confirm_preferences(
        self,
        ids: list,
        edits: Optional[dict] = None,
        reject: Optional[list] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Confirm proposed preferences ONLY after the user agreed to them when read back. Apply their corrections through edits; rejected ones through reject.

        :param ids: Preference ids the user confirmed.
        :param edits: {"<id>": {fields to change}} for corrections.
        :param reject: Preference ids the user rejected.
        :return: JSON with confirmed and rejected.
        """
        return await self._do(
            "POST",
            "/preferences/confirm",
            _clean({"ids": _coerce(ids), "edits": _coerce(edits), "reject": _coerce(reject)}),
            emitter=__event_emitter__,
            status="Confirming preferences",
        )

    # ------------------------------------------------------ household

    async def manage_household(
        self,
        action: str = "list",
        id: Optional[str] = None,
        label: Optional[str] = None,
        role: Optional[str] = None,
        birth_month: Optional[str] = None,
        allergies: Optional[list] = None,
        active: Optional[bool] = None,
        settings: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Household members and settings. Allergies are only ever added or removed on an explicit statement, read back to the user before saving.

        :param action: "list" (people), "save" (add or edit a person), "get_settings", or "set_settings".
        :param id: Person id when editing.
        :param label: Person's name or label.
        :param role: "adult" or "child".
        :param birth_month: YYYY-MM.
        :param allergies: List of allergens (the full list for that person).
        :param active: False to retire a person.
        :param settings: For set_settings: any of {portions:{adult,child}, child_cooldown_days, week_start_day, use_soon_days}.
        :return: JSON from the service.
        """
        a = (action or "list").lower()
        if a == "list":
            return await self._do("GET", "/people", emitter=__event_emitter__, status="Reading household")
        if a == "get_settings":
            return await self._do("GET", "/settings", emitter=__event_emitter__, status="Reading settings")
        if a == "set_settings":
            return await self._do(
                "PUT", "/settings", _coerce(settings) or {},
                emitter=__event_emitter__, status="Saving settings",
            )
        if a == "save":
            return await self._do(
                "POST",
                "/people",
                _clean(
                    {
                        "id": id,
                        "label": label,
                        "role": role,
                        "birth_month": birth_month,
                        "allergies": _coerce(allergies),
                        "active": active,
                    }
                ),
                emitter=__event_emitter__,
                status="Saving household member",
            )
        return _compact({"ok": False, "error": "invalid", "detail": "action must be list, save, get_settings or set_settings"})

    async def accuracy_report(
        self,
        since: Optional[str] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Evidence about where the pantry's believed stock drifts from reality (consistent drift, meals logged after the fact, never-restocked items, recipes that understate an ingredient). Turn each gap into a proposed fix in your own words; never apply a fix without the user.

        :param since: Optional YYYY-MM-DD start.
        :return: JSON with audits, per-item drift and gaps.
        """
        try:
            res = await self._call("GET", "/accuracy", None, {"since": since})
        except Exception as e:
            return _compact(self._error_result(e))
        return _compact(res)

    # ------------------------------------------------------ spreadsheets

    @staticmethod
    def _norm_header(h: Any) -> str:
        n = re.sub(r"[\s_]+", "_", str(h or "").strip().lower())
        return n if n in _CANONICAL else n.replace("_", " ")

    @staticmethod
    def _cell_text(v: Any) -> str:
        if v is None:
            return ""
        if isinstance(v, datetime.datetime):
            return v.date().isoformat() if v.time() == datetime.time(0) else v.isoformat()
        if isinstance(v, datetime.date):
            return v.isoformat()
        if isinstance(v, bool):
            return str(v)
        if isinstance(v, float) and v.is_integer() and abs(v) < 1e15:
            return _Num(int(v))
        if isinstance(v, (int, float)):
            return _Num(repr(v))
        return str(v).strip()

    @staticmethod
    def _decode(raw: bytes) -> str:
        for enc in ("utf-8-sig", "cp1252"):
            try:
                return raw.decode(enc)
            except UnicodeDecodeError:
                continue
        return raw.decode("latin-1")

    def _parse_table(self, raw: bytes, file_name: str) -> tuple:
        """Return (source, headers, rows) from RAW bytes. rows are lists of text cells."""
        name = (file_name or "").lower()
        is_xlsx = name.endswith((".xlsx", ".xlsm")) or raw[:2] == b"PK"
        if is_xlsx:
            import openpyxl  # shipped in the OWUI image

            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb[wb.sheetnames[0]]
            table = [[self._cell_text(c) for c in row] for row in ws.iter_rows(values_only=True)]
            wb.close()
            source = "xlsx"
        else:
            text = self._decode(raw)
            first = text.split("\n", 1)[0]
            delim = max([",", ";", "\t"], key=first.count) if first else ","
            table = [[c.strip() for c in r] for r in csv.reader(io.StringIO(text, newline=""), delimiter=delim)]
            source = "csv"
        table = [r for r in table if any(c != "" for c in r)]
        if not table:
            return source, [], []
        return source, table[0], table[1:]

    async def _read_attached(self, entry: dict, __user__: Optional[dict]) -> tuple:
        """Raw bytes + file name of one __files__ entry (never the RAG-extracted text)."""
        nested = entry.get("file") if isinstance(entry.get("file"), dict) else {}
        file_id = entry.get("id") or nested.get("id")
        name = (
            entry.get("name")
            or nested.get("filename")
            or (nested.get("meta") or {}).get("name")
            or ""
        )
        path = nested.get("path") or entry.get("path")
        if not path and file_id:
            from open_webui.models.files import Files  # only exists inside OWUI

            rec = await Files.get_file_by_id(file_id)
            if rec is None:
                raise FileNotFoundError(f"file {file_id} not found")
            path = rec.path
            name = name or rec.filename
        if not path:
            raise FileNotFoundError("attached file has no id or path")
        if not nested.get("path") and not entry.get("path"):
            from open_webui.storage.provider import Storage

            path = await asyncio.to_thread(Storage.get_file, path)
        def _rd():
            with open(path, "rb") as fh:
                return fh.read()

        raw = await asyncio.to_thread(_rd)
        return raw, name

    @staticmethod
    def _pick_file(files: Optional[list], file_name: Optional[str]) -> Optional[dict]:
        cands = []
        for f in files or []:
            if not isinstance(f, dict):
                continue
            nm = (f.get("name") or (f.get("file") or {}).get("filename") or "").lower()
            if nm.endswith((".csv", ".xlsx", ".xlsm", ".tsv", ".txt")):
                cands.append((nm, f))
        if file_name:
            for nm, f in cands:
                if nm == file_name.lower():
                    return f
        return cands[-1][1] if cands else None

    def _propose_mapping(self, headers: list) -> dict:
        proposed = {}
        for h in headers:
            n = self._norm_header(h)
            if n in _CANONICAL:
                proposed[h] = n
            elif n in _SYNONYMS:
                proposed[h] = _SYNONYMS[n]
            else:
                proposed[h] = None
        return proposed

    def _headers_are_exact(self, headers: list) -> bool:
        norm = [self._norm_header(h) for h in headers if str(h).strip()]
        return "name" in norm and all(n in _CANONICAL for n in norm) and len(set(norm)) == len(norm)

    @staticmethod
    def _to_number(text: str) -> Optional[float]:
        """A plain finite non-negative number, or None. Never guesses: "1,000", "1.000"
        (thousands or decimal?), "1,5", currency, units, nan/inf and negatives are all None.
        A genuinely numeric spreadsheet cell (_Num) is unambiguous, so "1.125" is accepted there."""
        t = str(text).strip()
        if not _PLAIN_NUMBER.fullmatch(t):
            return None
        if not isinstance(text, _Num) and _AMBIGUOUS_THOUSANDS.fullmatch(t):
            return None
        return float(t)

    def _build_rows(self, headers: list, rows: list, mapping: dict) -> tuple:
        """Map sheet rows to /audit/preview rows. Returns (rows, skipped_row_numbers, invalid)
        where invalid lists every quantity/actual cell that is not a plain number."""
        idx = {}
        for i, h in enumerate(headers):
            target = mapping.get(h)
            if target and target in _CANONICAL and target not in idx:
                idx[target] = i
        out, skipped, invalid = [], [], []
        for n, r in enumerate(rows, start=2):  # row 1 is the header
            def cell(field):
                i = idx.get(field)
                if i is None or i >= len(r):
                    return ""
                return r[i] if isinstance(r[i], _Num) else r[i].strip()

            row = {}
            for field in _ROW_FIELDS + _OPTIONAL_COLUMNS:
                val = cell(field)
                if val == "":
                    continue
                if field in _NUMERIC_FIELDS:
                    num = self._to_number(val)
                    if num is None:
                        invalid.append({"row": n, "column": field, "value": val})
                        continue
                    row[field] = int(num) if float(num).is_integer() else num
                elif field == "level":
                    if val.lower() in _LEVELS:
                        row[field] = val.lower()
                    else:
                        invalid.append({"row": n, "column": "level", "value": val})
                elif field in _LIST_FIELDS:
                    row[field] = [p.strip() for p in re.split(r"[;|,]", val) if p.strip()]
                else:
                    row[field] = val
            actual = cell("actual")
            if actual != "":
                kind = str(row.get("kind", "")).strip().lower()
                is_level = actual.lower() in _LEVELS
                if kind == "staple":
                    # a staple's actual is a level, exactly one of plenty|low|out
                    if is_level:
                        row["level"] = actual.lower()
                    else:
                        invalid.append({"row": n, "column": "actual", "value": actual})
                elif kind == "counted" and is_level:
                    # a counted item has a number, never a level: refuse, do not post both
                    invalid.append({"row": n, "column": "actual", "value": actual})
                elif is_level:  # kind unknown: a level word can only mean a level
                    row["level"] = actual.lower()
                else:
                    num = self._to_number(actual)
                    if num is None:
                        invalid.append({"row": n, "column": "actual", "value": actual})
                    else:
                        row["quantity"] = int(num) if float(num).is_integer() else num
            if "name" not in row and "id" not in row:
                skipped.append(n)
                continue
            if "name" not in row:
                row["name"] = ""
            out.append(row)
        return out, skipped, invalid

    async def import_pantry(
        self,
        file_name: Optional[str] = None,
        mapping: Optional[str] = None,
        __files__: Optional[list] = None,
        __user__: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Read the .csv or .xlsx the user attached to this chat (the RAW file, every row) and preview what would change in the pantry. Writes NOTHING. Show the diff to the user, then use confirm_import only after they agree.
        If the sheet's column names differ from the export's, this returns a proposed column mapping instead: show it to the user, and after they confirm call import_pantry again with mapping set.

        :param file_name: Which attached file, if several are attached.
        :param mapping: JSON object {"<sheet column>": "<pantry field or null>"} the user confirmed. Fields: id, name, category, kind, quantity, unit, level, location, expires_on, allergens, may_contain, actual.
        :return: JSON preview (preview_id, changed, new, missing, unmatched, unconvertible) or a mapping proposal.
        """
        entry = self._pick_file(__files__, file_name)
        if entry is None:
            return _compact(
                {"ok": False, "error": "no_file", "instruction": "No .csv or .xlsx is attached. Ask the user to attach the sheet."}
            )
        await self._emit(__event_emitter__, "Reading attached sheet", False)
        try:
            raw, name = await self._read_attached(entry, __user__)
            source, headers, rows = self._parse_table(raw, name)
        except Exception as e:
            return _compact(
                {"ok": False, "error": "unreadable_file", "detail": f"{type(e).__name__}: {e}",
                 "instruction": "Tell the user the file could not be read; do not guess its contents."}
            )
        if not headers:
            return _compact({"ok": False, "error": "empty_file", "instruction": "The sheet has no rows. Tell the user."})

        chosen = _coerce(mapping)
        if isinstance(chosen, str):
            chosen = None
        if chosen is None and not self._headers_are_exact(headers):
            proposal = self._propose_mapping(headers)
            return _compact(
                {
                    "ok": True,
                    "needs_confirmation": "column_mapping",
                    "file_name": name,
                    "row_count": len(rows),
                    "proposed_mapping": proposal,
                    "unmapped_columns": [h for h, t in proposal.items() if t is None],
                    "pantry_fields": _CANONICAL,
                    "sample_rows": [dict(zip(headers, r)) for r in rows[:3]],
                    "instruction": (
                        "Nothing was sent. Show this mapping to the user. After they confirm or correct it, "
                        "call import_pantry again with mapping set to the confirmed JSON object."
                    ),
                }
            )
        mp = chosen if chosen is not None else {h: self._norm_header(h) for h in headers}
        out_rows, skipped, invalid = self._build_rows(headers, rows, mp)
        if not any("name" in r or "id" in r for r in out_rows):
            return _compact({"ok": False, "error": "invalid", "detail": "no column is mapped to name or id", "proposed_mapping": self._propose_mapping(headers)})
        if invalid:
            return _compact(
                {"ok": False, "error": "invalid", "nothing_sent": True,
                 "detail": "some quantity/actual cells are not plain numbers, or some level/actual cells for staples are not plenty|low|out",
                 "invalid": invalid[:50], "invalid_count": len(invalid),
                 "instruction": (
                     "Nothing was sent. Show the user each row and value, and ask what number they mean "
                     "(for example 1,000 could be one thousand or one). Never fix a value yourself. "
                     "After they correct the sheet and attach it again, or tell you the values, call import_pantry again."
                 )}
            )
        body = {"source": source, "file_name": name, "rows": out_rows}
        try:
            res = await self._call("POST", "/audit/preview", body)
        except Exception as e:
            return _compact(self._error_result(e))
        await self._emit(__event_emitter__, "Preview ready", True)
        return _compact(self._shape_preview(res, len(out_rows), skipped))

    def _shape_preview(self, res: Any, sent: int, skipped: list) -> dict:
        if not isinstance(res, dict):
            return {"ok": True, "preview": res}
        cap = max(1, int(self.valves.preview_max_lines))
        out = {"ok": True, "rows_sent": sent}
        if skipped:
            out["skipped_rows_without_name_or_id"] = skipped[:50]
        for k, v in res.items():
            if isinstance(v, list):
                out[k] = v[:cap]
                out[f"{k}_count"] = len(v)
                if len(v) > cap:
                    out[f"{k}_truncated"] = True
            else:
                out[k] = v
        out["instruction"] = (
            "Show the user this diff (changed, new, missing). Items in missing are asked about, never "
            "zeroed on your own. unmatched/unconvertible need the user. Commit only with confirm_import "
            "after they agree."
        )
        return out

    async def confirm_import(
        self,
        preview_id: str,
        missing: Optional[dict] = None,
        exclude: Optional[list] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Commit a previewed import AFTER the user confirmed the diff. Items missing from the sheet default to keep unless the user decided otherwise.

        :param preview_id: The preview_id from import_pantry.
        :param missing: {"<item id>": "keep" | "zero" | "remove"} for items missing from the sheet.
        :param exclude: Item ids to leave out of the commit.
        :return: JSON with audit_id, items_checked, items_changed.
        """
        return await self._do(
            "POST",
            "/audit/commit",
            _clean({"preview_id": preview_id, "missing": _coerce(missing), "exclude": _coerce(exclude)}),
            emitter=__event_emitter__,
            status="Committing audit",
        )

    # ------------------------------------------------------------ export

    def _export_rows(self, items: list) -> list:
        rows = []
        for it in items:
            row = []
            for col in EXPORT_COLUMNS:
                if col == "actual":
                    row.append("")
                    continue
                v = it.get(col)
                if isinstance(v, list):
                    v = ";".join(str(x) for x in v)
                elif v is None:
                    v = ""
                elif isinstance(v, float) and v.is_integer():
                    v = int(v)
                row.append(v)
            rows.append(row)
        return rows

    def _build_csv(self, items: list) -> bytes:
        buf = io.StringIO(newline="")
        w = csv.writer(buf, lineterminator="\n")
        w.writerow(EXPORT_COLUMNS)
        w.writerows(self._export_rows(items))
        return ("﻿" + buf.getvalue()).encode("utf-8")

    def _build_xlsx(self, items: list) -> bytes:
        import openpyxl

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "pantry"
        ws.append(EXPORT_COLUMNS)
        for r in self._export_rows(items):
            ws.append(r)
        out = io.BytesIO()
        wb.save(out)
        return out.getvalue()

    async def _deliver_file(self, name: str, data: bytes, content_type: str, user: Optional[dict], emitter) -> Optional[dict]:
        """Store the bytes in OWUI's Files store and attach them to the chat message.
        Returns {"id","url"} or None when the OWUI internals are unavailable."""
        try:
            from open_webui.models.files import FileForm, Files
            from open_webui.storage.provider import Storage

            uid = (user or {}).get("id")
            if not uid:
                return None
            fid = str(uuid.uuid4())
            tags = {"OpenWebUI-User-Id": str(uid), "OpenWebUI-File-Id": fid}
            _, path = await asyncio.to_thread(Storage.upload_file, io.BytesIO(data), f"{fid}_{name}", tags)
            rec = Files.insert_new_file(
                uid,
                FileForm(
                    id=fid,
                    filename=name,
                    path=path,
                    data={},
                    meta={"name": name, "content_type": content_type, "size": len(data)},
                ),
            )
            if asyncio.iscoroutine(rec):
                rec = await rec
            if rec is None:
                return None
            url = f"/api/v1/files/{fid}/content"
            if emitter:
                await emitter({"type": "files", "data": {"files": [{"type": "file", "url": url, "name": name}]}})
            return {"id": fid, "url": url}
        except Exception:
            return None

    async def export_pantry(
        self,
        format: str = "csv",
        __user__: Optional[dict] = None,
        __event_emitter__: Optional[Callable[[dict], Awaitable[None]]] = None,
    ) -> str:
        """
        Export the whole pantry as a spreadsheet the user can edit and re-import. CSV is the default; use "xlsx" only if asked.
        Columns: id, name, category, kind, quantity, unit, level, location, expires_on, allergens, actual (empty; the user fills in what is really there).

        :param format: "csv" (default) or "xlsx".
        :return: A download link for the file, or the CSV inline in a fenced block if a download is not possible.
        """
        fmt = (format or "csv").lower().strip()
        if fmt not in ("csv", "xlsx"):
            return _compact({"ok": False, "error": "invalid", "detail": 'format must be "csv" or "xlsx"'})
        await self._emit(__event_emitter__, "Exporting pantry", False)
        try:
            res = await self._call("GET", "/pantry")
        except Exception as e:
            return _compact(self._error_result(e))
        items = res.get("items", []) if isinstance(res, dict) else []
        stamp = datetime.date.today().isoformat()
        csv_bytes = self._build_csv(items)
        if fmt == "xlsx":
            try:
                data, ctype, fname = (
                    self._build_xlsx(items),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    f"pantry-{stamp}.xlsx",
                )
            except ImportError:
                data, ctype, fname, fmt = csv_bytes, "text/csv", f"pantry-{stamp}.csv", "csv"
        else:
            data, ctype, fname = csv_bytes, "text/csv", f"pantry-{stamp}.csv"
        delivered = await self._deliver_file(fname, data, ctype, __user__, __event_emitter__)
        await self._emit(__event_emitter__, "Export ready", True)
        if delivered:
            return (
                f"Exported {len(items)} items as {fname}. Tell the user to download it here: "
                f"[{fname}]({delivered['url']}?attachment=true). Edit the `actual` column (or any "
                f"column), then attach the file back to import it."
            )
        text = csv_bytes.decode("utf-8").lstrip("﻿")
        return (
            f"Download is not available here, so here is the pantry ({len(items)} items) as CSV. "
            "Tell the user to copy it into a .csv file:\n\n```csv\n" + text + "```"
        )
